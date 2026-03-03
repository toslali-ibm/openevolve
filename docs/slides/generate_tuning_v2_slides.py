#!/usr/bin/env python3
"""Generate 4 slide variants introducing the Tuning v2 (tiered rescue/polish) mechanism."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

# ─── Colors ───────────────────────────────────────────────────────────────────

DARK_BLUE = RGBColor(0x00, 0x2D, 0x5E)
MID_BLUE = RGBColor(0x00, 0x53, 0xA0)
LIGHT_BLUE = RGBColor(0x41, 0x78, 0xBE)
ACCENT_TEAL = RGBColor(0x00, 0xB3, 0x9E)
ACCENT_GREEN = RGBColor(0x24, 0xA1, 0x48)
ACCENT_RED = RGBColor(0xDA, 0x1E, 0x28)
ACCENT_ORANGE = RGBColor(0xFF, 0x83, 0x2B)
ACCENT_PURPLE = RGBColor(0x8A, 0x3F, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF4, 0xF4, 0xF4)
MED_GRAY = RGBColor(0x8D, 0x8D, 0x8D)
DARK_GRAY = RGBColor(0x39, 0x39, 0x39)
BLACK = RGBColor(0x00, 0x00, 0x00)
CODE_BG = RGBColor(0x26, 0x26, 0x2E)
SOFT_YELLOW = RGBColor(0xFF, 0xF3, 0xCD)
SOFT_GREEN = RGBColor(0xD4, 0xED, 0xDA)
SOFT_RED = RGBColor(0xF8, 0xD7, 0xDA)
SOFT_BLUE = RGBColor(0xD6, 0xEA, 0xF8)


# ─── Helpers ──────────────────────────────────────────────────────────────────

def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=BLACK, bold=False, alignment=PP_ALIGN.LEFT,
                font_name="Arial"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_multiline(slide, left, top, width, height, lines,
                  font_size=16, color=BLACK, font_name="Arial",
                  line_spacing=None):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if isinstance(line, tuple):
            text, is_bold, clr = line
        else:
            text, is_bold, clr = line, False, color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(font_size)
        p.font.color.rgb = clr
        p.font.bold = is_bold
        p.font.name = font_name
        p.space_after = Pt(line_spacing if line_spacing else font_size * 0.3)
    return txBox


def add_box(slide, left, top, width, height, text, fill_color,
            text_color=WHITE, font_size=14, bold=True, font_name="Arial",
            alignment=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(8)
    tf.margin_right = Pt(8)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = text_color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    tf.auto_size = None
    return shape


def add_multiline_box(slide, left, top, width, height, lines, fill_color,
                      text_color=WHITE, font_size=13, font_name="Arial"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(10)
    tf.margin_right = Pt(10)
    tf.margin_top = Pt(6)
    tf.margin_bottom = Pt(6)
    for i, line in enumerate(lines):
        if isinstance(line, tuple):
            text, is_bold, clr, sz = line
        else:
            text, is_bold, clr, sz = line, False, text_color, font_size
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(sz)
        p.font.color.rgb = clr
        p.font.bold = is_bold
        p.font.name = font_name
        p.space_after = Pt(2)
    tf.auto_size = None
    return shape


def add_arrow(slide, left, top, width, height, color=LIGHT_BLUE,
              direction="down"):
    shape_type = {
        "down": MSO_SHAPE.DOWN_ARROW,
        "right": MSO_SHAPE.RIGHT_ARROW,
    }.get(direction, MSO_SHAPE.DOWN_ARROW)
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_circle(slide, left, top, size, text, fill_color, text_color=WHITE,
               font_size=14, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = text_color
    p.font.bold = bold
    p.font.name = "Arial"
    p.alignment = PP_ALIGN.CENTER
    tf.auto_size = None
    return shape


def add_diamond(slide, left, top, width, height, text, fill_color,
                text_color=WHITE, font_size=12, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = text_color
    p.font.bold = bold
    p.font.name = "Arial"
    p.alignment = PP_ALIGN.CENTER
    tf.auto_size = None
    return shape


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 1: "The Three Tiers" — staircase showing rescue / checkpoint / final
# ═══════════════════════════════════════════════════════════════════════════════

def create_slide_1():
    prs = new_prs()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)

    # Title
    add_textbox(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7),
                "Tuning v2: Three Tiers of Optimization",
                font_size=32, color=DARK_BLUE, bold=True)
    add_textbox(slide, Inches(0.6), Inches(0.9), Inches(12), Inches(0.5),
                "Graduate the tuning budget by context: cheap screening, moderate polish, thorough final pass",
                font_size=18, color=MED_GRAY)

    # ── Three tier columns ──
    col_w = Inches(3.6)
    gap = Inches(0.4)

    # Tier 1: Rescue
    c1 = Inches(0.6)
    add_box(slide, c1, Inches(1.8), col_w, Inches(0.55),
            "Tier 1: Rescue", ACCENT_ORANGE, font_size=18)

    add_multiline_box(slide, c1, Inches(2.5), col_w, Inches(1.6), [
        ("5 trials", True, WHITE, 20),
        ("", False, WHITE, 6),
        ("When: Every iteration (selective)", False, WHITE, 13),
        ("Who: Novel programs scoring poorly", False, WHITE, 13),
        ("Why: Is this algorithm salvageable?", False, WHITE, 13),
    ], ACCENT_ORANGE)

    # Budget bar for Tier 1
    bar_y = Inches(4.35)
    bar_h = Inches(0.5)
    add_box(slide, c1, bar_y, Inches(0.9), bar_h,
            "5", ACCENT_ORANGE, font_size=16)
    add_textbox(slide, c1 + Inches(1.0), bar_y + Inches(0.1), Inches(2.5), Inches(0.4),
                "trials per program", font_size=13, color=DARK_GRAY)

    # Tier 2: Checkpoint Polish
    c2 = c1 + col_w + gap
    add_box(slide, c2, Inches(1.8), col_w, Inches(0.55),
            "Tier 2: Checkpoint Polish", ACCENT_TEAL, font_size=18)

    add_multiline_box(slide, c2, Inches(2.5), col_w, Inches(1.6), [
        ("10 trials", True, WHITE, 20),
        ("", False, WHITE, 6),
        ("When: At each checkpoint interval", False, WHITE, 13),
        ("Who: Top-3 elites per island", False, WHITE, 13),
        ("Why: Improve parents for next gen", False, WHITE, 13),
    ], ACCENT_TEAL)

    bar_w2 = Inches(1.8)
    add_box(slide, c2, bar_y, bar_w2, bar_h,
            "10", ACCENT_TEAL, font_size=16)
    add_textbox(slide, c2 + bar_w2 + Inches(0.1), bar_y + Inches(0.1), Inches(2), Inches(0.4),
                "trials per program", font_size=13, color=DARK_GRAY)

    # Tier 3: Final Polish
    c3 = c2 + col_w + gap
    add_box(slide, c3, Inches(1.8), col_w, Inches(0.55),
            "Tier 3: Final Polish", MID_BLUE, font_size=18)

    add_multiline_box(slide, c3, Inches(2.5), col_w, Inches(1.6), [
        ("20 trials", True, WHITE, 20),
        ("", False, WHITE, 6),
        ("When: Final iteration only", False, WHITE, 13),
        ("Who: Top-3 elites per island", False, WHITE, 13),
        ("Why: Thorough optimization at the end", False, WHITE, 13),
    ], MID_BLUE)

    bar_w3 = Inches(3.6)
    add_box(slide, c3, bar_y, bar_w3, bar_h,
            "20", MID_BLUE, font_size=16)

    # ── Bottom: 25-iteration timeline ──
    add_textbox(slide, Inches(0.6), Inches(5.2), Inches(12), Inches(0.5),
                "Example: 25-iteration run (checkpoint every 5 iterations)",
                font_size=16, color=DARK_BLUE, bold=True)

    tl_y = Inches(5.8)
    tl_left = Inches(0.8)
    iter_w = Inches(0.42)

    for i in range(25):
        x = tl_left + i * (iter_w + Inches(0.05))
        is_checkpoint = (i + 1) % 5 == 0 and (i + 1) < 25
        is_final = (i + 1) == 25

        if is_final:
            color = MID_BLUE
        elif is_checkpoint:
            color = ACCENT_TEAL
        else:
            color = LIGHT_GRAY

        add_box(slide, x, tl_y, iter_w, Inches(0.4),
                str(i + 1), color,
                text_color=WHITE if color != LIGHT_GRAY else DARK_GRAY,
                font_size=9, bold=False)

    # Rescue markers (scattered ~32% of iterations)
    rescue_iters = [1, 3, 7, 10, 14, 18, 21, 24]
    for ri in rescue_iters:
        x = tl_left + (ri - 1) * (iter_w + Inches(0.05))
        add_box(slide, x, tl_y + Inches(0.45), iter_w, Inches(0.15),
                "", ACCENT_ORANGE, font_size=6)

    # Legend
    lg_y = Inches(6.6)
    add_box(slide, Inches(1.0), lg_y, Inches(0.4), Inches(0.3),
            "", ACCENT_ORANGE, font_size=8)
    add_textbox(slide, Inches(1.5), lg_y, Inches(2.2), Inches(0.3),
                "Rescue (~32% of iterations)", font_size=12, color=DARK_GRAY)

    add_box(slide, Inches(4.3), lg_y, Inches(0.4), Inches(0.3),
            "", ACCENT_TEAL, font_size=8)
    add_textbox(slide, Inches(4.8), lg_y, Inches(2.2), Inches(0.3),
                "Checkpoint polish (iter 5,10,15,20)", font_size=12, color=DARK_GRAY)

    add_box(slide, Inches(8.0), lg_y, Inches(0.4), Inches(0.3),
            "", MID_BLUE, font_size=8)
    add_textbox(slide, Inches(8.5), lg_y, Inches(2.2), Inches(0.3),
                "Final polish (iter 25)", font_size=12, color=DARK_GRAY)

    # Total budget
    add_textbox(slide, Inches(0.6), Inches(7.0), Inches(12), Inches(0.4),
                "Total: ~220 evaluations (vs v1's 500) — 56% fewer, graduated by value",
                font_size=14, color=MID_BLUE, bold=True, alignment=PP_ALIGN.CENTER)

    prs.save(os.path.join(OUT_DIR, "tuning_v2_three_tiers.pptx"))
    print("Created: tuning_v2_three_tiers.pptx")


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 2: "Smart Selection" — decision flowchart for rescue
# ═══════════════════════════════════════════════════════════════════════════════

def create_slide_2():
    prs = new_prs()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)

    # Title
    add_textbox(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7),
                "Smart Rescue: Only Tune Where It Matters",
                font_size=32, color=DARK_BLUE, bold=True)
    add_textbox(slide, Inches(0.6), Inches(0.9), Inches(12), Inches(0.5),
                "v1 tuned every program (wasteful). v2 only tunes novel algorithms with bad numbers.",
                font_size=18, color=MED_GRAY)

    # ── Left side: Decision flowchart ──
    add_textbox(slide, Inches(0.6), Inches(1.7), Inches(6), Inches(0.5),
                "Per-Iteration Decision Flow", font_size=20, color=DARK_BLUE, bold=True)

    fx = Inches(1.5)

    # Step 1
    add_box(slide, fx, Inches(2.3), Inches(4.5), Inches(0.5),
            "LLM generates code", MID_BLUE, font_size=14)

    add_arrow(slide, fx + Inches(2.0), Inches(2.85), Inches(0.4), Inches(0.3), MED_GRAY)

    # Decision: has @TUNE?
    add_diamond(slide, fx + Inches(0.7), Inches(3.2), Inches(3), Inches(1.0),
                "Has @TUNE?", LIGHT_BLUE, font_size=13)

    # No path (right)
    add_arrow(slide, fx + Inches(3.7), Inches(3.5), Inches(0.6), Inches(0.3),
              MED_GRAY, direction="right")
    add_box(slide, fx + Inches(4.4), Inches(3.4), Inches(1.8), Inches(0.45),
            "Skip (no cost)", MED_GRAY, font_size=11)
    add_textbox(slide, fx + Inches(3.8), Inches(3.15), Inches(0.6), Inches(0.3),
                "No", font_size=11, color=ACCENT_RED, bold=True)

    # Yes path (down)
    add_arrow(slide, fx + Inches(1.9), Inches(4.2), Inches(0.4), Inches(0.3), ACCENT_TEAL)
    add_textbox(slide, fx + Inches(1.2), Inches(4.2), Inches(0.6), Inches(0.3),
                "Yes", font_size=11, color=ACCENT_GREEN, bold=True)

    # Evaluate pre-tune
    add_box(slide, fx, Inches(4.55), Inches(4.5), Inches(0.45),
            "Evaluate program (get pre-tune score)", LIGHT_BLUE, font_size=13)

    add_arrow(slide, fx + Inches(2.0), Inches(5.05), Inches(0.4), Inches(0.3), ACCENT_TEAL)

    # Decision: rescue criteria
    add_diamond(slide, fx + Inches(0.2), Inches(5.4), Inches(3.8), Inches(1.1),
                "Score < median\nAND novel?", ACCENT_ORANGE, font_size=11)

    # Yes path -> rescue
    add_textbox(slide, fx - Inches(0.3), Inches(5.75), Inches(0.6), Inches(0.3),
                "Yes", font_size=11, color=ACCENT_GREEN, bold=True)
    add_arrow(slide, fx + Inches(2.0), Inches(6.5), Inches(0.4), Inches(0.3), ACCENT_GREEN)
    add_box(slide, fx + Inches(0.2), Inches(6.85), Inches(3.8), Inches(0.45),
            "RESCUE: Tune with 5 trials", ACCENT_GREEN, font_size=13)

    # No path -> skip
    add_arrow(slide, fx + Inches(4.0), Inches(5.75), Inches(0.6), Inches(0.3),
              MED_GRAY, direction="right")
    add_textbox(slide, fx + Inches(4.1), Inches(5.45), Inches(0.5), Inches(0.3),
                "No", font_size=11, color=ACCENT_RED, bold=True)
    add_box(slide, fx + Inches(4.7), Inches(5.65), Inches(1.8), Inches(0.45),
            "Skip (reuse eval)", MED_GRAY, font_size=11)

    # ── Right side: Why selective? ──
    rx = Inches(7.5)
    add_textbox(slide, rx, Inches(1.7), Inches(5.5), Inches(0.5),
                "Why Be Selective?", font_size=20, color=DARK_BLUE, bold=True)

    # v1 stats box
    add_multiline_box(slide, rx, Inches(2.3), Inches(5.2), Inches(1.7), [
        ("v1: Tuned every program, every iteration", True, WHITE, 14),
        ("", False, WHITE, 4),
        ("68% of tuning was on already-good programs", False, WHITE, 13),
        ("Average gain on good programs: +0.01", False, WHITE, 13),
        ("Average gain on novel-but-poor: +0.47", False, WHITE, 13),
        ("Tuning good programs = 47x less efficient", False, SOFT_YELLOW, 13),
    ], ACCENT_RED)

    # v2 stats box
    add_multiline_box(slide, rx, Inches(4.2), Inches(5.2), Inches(1.7), [
        ("v2: Tune only where ROI is high", True, WHITE, 14),
        ("", False, WHITE, 4),
        ("~32% of programs get rescue tuning", False, WHITE, 13),
        ("5 trials per rescue (vs v1's 19)", False, WHITE, 13),
        ("Pre-eval is reused (zero redundant evals)", False, WHITE, 13),
        ("56% fewer total evaluations", False, SOFT_GREEN, 13),
    ], ACCENT_GREEN)

    # Key insight
    add_box(slide, rx, Inches(6.2), Inches(5.2), Inches(0.9),
            "Key Insight\nNovel algorithm + bad numbers = rescue candidate\n"
            "Good score OR unoriginal = skip",
            DARK_BLUE, font_size=13, bold=False)

    prs.save(os.path.join(OUT_DIR, "tuning_v2_smart_selection.pptx"))
    print("Created: tuning_v2_smart_selection.pptx")


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3: "v1 vs v2" — side-by-side comparison
# ═══════════════════════════════════════════════════════════════════════════════

def create_slide_3():
    prs = new_prs()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)

    # Title
    add_textbox(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7),
                "Tuning v1 vs v2: What Changed and Why",
                font_size=32, color=DARK_BLUE, bold=True)
    add_textbox(slide, Inches(0.6), Inches(0.9), Inches(12), Inches(0.5),
                "v1 steered the LLM toward parameters. v2 keeps the LLM focused on algorithms.",
                font_size=18, color=MED_GRAY)

    # Column headers
    row_h = Inches(0.8)
    label_w = Inches(2.5)
    col_w = Inches(4.8)
    lx = Inches(0.6)
    v1x = lx + label_w + Inches(0.15)
    v2x = v1x + col_w + Inches(0.15)

    header_y = Inches(1.7)
    add_box(slide, lx, header_y, label_w, Inches(0.5),
            "", DARK_BLUE, font_size=14)
    add_box(slide, v1x, header_y, col_w, Inches(0.5),
            "v1 (tune everything)", ACCENT_RED, font_size=15)
    add_box(slide, v2x, header_y, col_w, Inches(0.5),
            "v2 (smart + tiered)", ACCENT_GREEN, font_size=15)

    # Rows
    rows = [
        ("LLM Prompt",
         "30 lines, 12x @TUNE mentions\nFrames LLM as parameter selector",
         "4 lines, 2x @TUNE mentions\nFrames LLM as algorithm inventor"),
        ("@TUNED Feedback",
         "Shown to LLM\nAnchors LLM on parameter thinking",
         "Hidden from LLM\nLLM only sees optimized values"),
        ("When to Tune",
         "Every program, every iteration\n68% of effort wasted on good programs",
         "Selective rescue (~32% of iterations)\n+ polish at checkpoints and final"),
        ("Trial Budget",
         "19 trials per program (flat)\n+ 1 redundant post-tune eval",
         "5 rescue / 10 checkpoint / 20 final\nZero redundant evaluations"),
        ("Total Evals\n(25 iterations)",
         "~500 evaluations\n~2.8 hours for BLIS",
         "~220 evaluations\n~1.2 hours for BLIS"),
        ("LLM Creativity",
         "LLM designs for tunability\nStructural innovation = rare",
         "LLM focuses on algorithms\nSame innovation rate as control"),
    ]

    for i, (label, v1_text, v2_text) in enumerate(rows):
        y = Inches(2.35) + i * (row_h + Inches(0.08))
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE

        add_multiline_box(slide, lx, y, label_w, row_h, [
            (label, True, DARK_BLUE, 12),
        ], bg, text_color=DARK_BLUE)

        # v1 cell
        v1_lines = v1_text.split("\n")
        add_multiline_box(slide, v1x, y, col_w, row_h, [
            (v1_lines[0], False, DARK_GRAY, 12),
        ] + [(l, False, MED_GRAY, 11) for l in v1_lines[1:]],
            SOFT_RED, text_color=DARK_GRAY)

        # v2 cell
        v2_lines = v2_text.split("\n")
        add_multiline_box(slide, v2x, y, col_w, row_h, [
            (v2_lines[0], True, DARK_GRAY, 12),
        ] + [(l, False, MED_GRAY, 11) for l in v2_lines[1:]],
            SOFT_GREEN, text_color=DARK_GRAY)

    # Bottom takeaway
    add_box(slide, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.4),
            "Core principle: the LLM invents algorithms, the optimizer handles numbers. "
            "Less prompt about tuning = more LLM creativity.",
            DARK_BLUE, font_size=13, bold=False)

    prs.save(os.path.join(OUT_DIR, "tuning_v2_comparison.pptx"))
    print("Created: tuning_v2_comparison.pptx")


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4: "End-to-End" — what happens to one program across the pipeline
# ═══════════════════════════════════════════════════════════════════════════════

def create_slide_4():
    prs = new_prs()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)

    # Title
    add_textbox(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7),
                "A Program's Journey Through Tuning v2",
                font_size=32, color=DARK_BLUE, bold=True)
    add_textbox(slide, Inches(0.6), Inches(0.9), Inches(12), Inches(0.5),
                "Follow one program from LLM generation through rescue, evolution, and final polish",
                font_size=18, color=MED_GRAY)

    # ── Top row: the rescue path ──
    add_textbox(slide, Inches(0.6), Inches(1.7), Inches(4), Inches(0.4),
                "Rescue Path (per iteration)", font_size=18, color=ACCENT_ORANGE, bold=True)

    bh = Inches(0.7)
    sw = Inches(2.8)  # step width
    ax = Inches(0.25)  # arrow width space

    # Step 1: LLM generates
    s1x = Inches(0.6)
    add_multiline_box(slide, s1x, Inches(2.2), sw, bh + Inches(0.3), [
        ("LLM Generates", True, WHITE, 14),
        ("temp = 0.5  # @TUNE [0.1, 2.0]", False, RGBColor(0xE0, 0xE0, 0xE0), 10),
        ("decay = 0.01 # @TUNE [0.001, 0.1]", False, RGBColor(0xE0, 0xE0, 0xE0), 10),
    ], MID_BLUE)

    add_arrow(slide, s1x + sw + Inches(0.05), Inches(2.55),
              Inches(0.4), Inches(0.3), MED_GRAY, direction="right")

    # Step 2: Pre-evaluate
    s2x = s1x + sw + Inches(0.55)
    add_multiline_box(slide, s2x, Inches(2.2), sw, bh + Inches(0.3), [
        ("Pre-Evaluate", True, WHITE, 14),
        ("Score: 0.38", False, WHITE, 12),
        ("Below median (0.72)", False, SOFT_YELLOW, 11),
    ], LIGHT_BLUE)

    add_arrow(slide, s2x + sw + Inches(0.05), Inches(2.55),
              Inches(0.4), Inches(0.3), MED_GRAY, direction="right")

    # Step 3: Rescue decision
    s3x = s2x + sw + Inches(0.55)
    add_multiline_box(slide, s3x, Inches(2.2), sw, bh + Inches(0.3), [
        ("Rescue Decision", True, WHITE, 14),
        ("Novel? YES (high diversity)", False, SOFT_GREEN, 11),
        ("Low score? YES (< median)", False, SOFT_GREEN, 11),
    ], ACCENT_ORANGE)

    add_arrow(slide, s3x + sw + Inches(0.05), Inches(2.55),
              Inches(0.4), Inches(0.3), MED_GRAY, direction="right")

    # Step 4: Optuna rescue
    s4x = s3x + sw + Inches(0.55)
    add_multiline_box(slide, s4x, Inches(2.2), sw, bh + Inches(0.3), [
        ("Optuna Rescue (5 trials)", True, WHITE, 14),
        ("temp: 0.5 -> 1.2", False, SOFT_GREEN, 11),
        ("Score: 0.38 -> 0.89", False, SOFT_GREEN, 11),
    ], ACCENT_GREEN)

    # ── Middle: what the LLM sees vs doesn't see ──
    mid_y = Inches(3.6)
    add_textbox(slide, Inches(0.6), mid_y, Inches(6), Inches(0.4),
                "What the LLM Sees (minimal prompt, no @TUNED)", font_size=18,
                color=DARK_BLUE, bold=True)

    add_multiline_box(slide, Inches(0.6), mid_y + Inches(0.45), Inches(5.5), Inches(1.2), [
        ("## Automatic Parameter Optimization", True, RGBColor(0xA0, 0xD0, 0xFF), 11),
        ("", False, WHITE, 4),
        ("An optimizer automatically tunes numeric parameters", False, RGBColor(0xE0, 0xE0, 0xE0), 11),
        ("in your code after generation. If your code has a", False, RGBColor(0xE0, 0xE0, 0xE0), 11),
        ("numeric value you're uncertain about, mark it:", False, RGBColor(0xE0, 0xE0, 0xE0), 11),
        ("  var = value  # @TUNE [min, max]", False, RGBColor(0x80, 0xFF, 0x80), 11),
        ("Focus on better algorithms, not better numbers.", False, RGBColor(0xE0, 0xE0, 0xE0), 11),
    ], CODE_BG)

    add_textbox(slide, Inches(7), mid_y, Inches(6), Inches(0.4),
                "What the LLM Does NOT See", font_size=18,
                color=ACCENT_RED, bold=True)

    add_multiline_box(slide, Inches(7), mid_y + Inches(0.45), Inches(5.5), Inches(1.2), [
        ("@TUNED annotations (gain, was, best_impact)", False, ACCENT_RED, 13),
        ("Tuning trial results or metrics", False, ACCENT_RED, 13),
        ("Which parameters were sensitive", False, ACCENT_RED, 13),
        ("Any guidance on parameter selection", False, ACCENT_RED, 13),
        ("", False, WHITE, 6),
        ("LLM stays focused on algorithms, not numbers", True, DARK_GRAY, 13),
    ], SOFT_RED, text_color=DARK_GRAY)

    # ── Bottom: Polish path ──
    bot_y = Inches(5.5)
    add_textbox(slide, Inches(0.6), bot_y, Inches(6), Inches(0.4),
                "Polish Path (at checkpoints and final iteration)",
                font_size=18, color=ACCENT_TEAL, bold=True)

    pw = Inches(3.5)

    # Checkpoint polish
    add_multiline_box(slide, Inches(0.6), bot_y + Inches(0.5), pw, Inches(1.3), [
        ("Checkpoint Polish", True, WHITE, 15),
        ("", False, WHITE, 4),
        ("Every 5 iterations", False, WHITE, 12),
        ("Top-3 elites per island", False, WHITE, 12),
        ("10 trials each", False, WHITE, 12),
        ("Improves parents for next gen", False, SOFT_GREEN, 11),
    ], ACCENT_TEAL)

    add_arrow(slide, Inches(4.3), bot_y + Inches(0.95),
              Inches(0.5), Inches(0.3), MED_GRAY, direction="right")

    # Iterations continue
    add_multiline_box(slide, Inches(5.0), bot_y + Inches(0.5), Inches(2.5), Inches(1.3), [
        ("Evolution Continues", True, WHITE, 15),
        ("", False, WHITE, 4),
        ("LLM builds on polished", False, WHITE, 12),
        ("parents with better", False, WHITE, 12),
        ("parameter values", False, WHITE, 12),
    ], MED_GRAY)

    add_arrow(slide, Inches(7.7), bot_y + Inches(0.95),
              Inches(0.5), Inches(0.3), MED_GRAY, direction="right")

    # Final polish
    add_multiline_box(slide, Inches(8.4), bot_y + Inches(0.5), pw, Inches(1.3), [
        ("Final Polish", True, WHITE, 15),
        ("", False, WHITE, 4),
        ("Last iteration only", False, WHITE, 12),
        ("Top-3 elites per island", False, WHITE, 12),
        ("20 trials each (thorough)", False, WHITE, 12),
        ("Squeeze out final parameter gains", False, SOFT_GREEN, 11),
    ], MID_BLUE)

    prs.save(os.path.join(OUT_DIR, "tuning_v2_journey.pptx"))
    print("Created: tuning_v2_journey.pptx")


# ─── Main ─────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    create_slide_1()
    create_slide_2()
    create_slide_3()
    create_slide_4()
    print("\nAll 4 tuning v2 slides generated in:", OUT_DIR)
